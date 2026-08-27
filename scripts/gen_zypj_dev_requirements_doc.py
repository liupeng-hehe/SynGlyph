# -*- coding: utf-8 -*-
"""Generate lean dev requirements Word doc with flowchart images."""

import os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = r"C:\Users\Administrator\Desktop\执与自组配件出入库-开发需求说明.docx"
ASSET_DIR = r"C:\Users\Administrator\Desktop\_zypj_flowcharts"
DOC_VERSION = "V2.0（开发版）"


def get_font(size=16, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    return ImageFont.load_default()


def draw_box(draw, xy, text, font, fill, outline="#333333", radius=8):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=2)
    lines = text.split("\n")
    lh = font.size + 6
    total_h = lh * len(lines)
    cy = (y1 + y2 - total_h) / 2
    for i, line in enumerate(lines):
        bbox = draw.textbbox((0, 0), line, font=font)
        tw = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - tw) / 2, cy + i * lh), line, fill="#111111", font=font)


def arrow_down(draw, x, y1, y2, color="#444444"):
    draw.line([(x, y1), (x, y2 - 10)], fill=color, width=2)
    draw.polygon([(x, y2), (x - 7, y2 - 12), (x + 7, y2 - 12)], fill=color)


def arrow_right(draw, x1, y, x2, color="#444444"):
    draw.line([(x1, y), (x2 - 10, y)], fill=color, width=2)
    draw.polygon([(x2, y), (x2 - 12, y - 7), (x2 - 12, y + 7)], fill=color)


def save_flowchart_overall(path):
    w, h = 1100, 1500
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    f = get_font(18)
    fb = get_font(20, bold=True)
    fs = get_font(16)

    draw.text((w // 2 - 200, 20), "图1  总流程：采购 → 组合 → 发货", fill="#111", font=fb)

    y = 70
    draw.text((40, y), "① 采购入库（现有流程，不变）", fill="#1565C0", font=fb)
    y += 40
    boxes = ["采购合同", "散件到货", "散件入【主库】"]
    bx = 120
    for i, t in enumerate(boxes):
        draw_box(draw, (bx, y, bx + 220, y + 56), t, f, "#E3F2FD")
        if i < len(boxes) - 1:
            arrow_right(draw, bx + 220, y + 28, bx + 300)
        bx += 300
    y += 90
    arrow_down(draw, 270, y - 34, y + 10)
    y += 20

    draw.text((40, y), "② 组合备货（新增，商务操作）", fill="#1565C0", font=fb)
    y += 40
    asm = [
        "商务维护BOM",
        "选组合包+数量",
        "系统算扣料",
        "审核",
        "散件出主库",
        "组合包入中间库",
    ]
    for i, t in enumerate(asm):
        draw_box(draw, (120, y, 420, y + 50), t, fs, "#FFF3E0")
        if i < len(asm) - 1:
            arrow_down(draw, 270, y + 50, y + 68)
        y += 68
    y += 10
    draw_box(draw, (500, y - 68 * 3 - 50, 780, y - 68 * 3), "主库散件\n够不够?", fs, "#FFEBEE")
    draw.text((790, y - 68 * 3 + 10), "不够→禁止", fill="#C62828", font=fs)
    y += 10

    draw.text((40, y), "③ 销售发货（发货时选产品类型）", fill="#1565C0", font=fb)
    y += 40
    draw_box(draw, (120, y, 400, y + 50), "销售合同 / 发货单", f, "#E8F5E9")
    arrow_down(draw, 260, y + 50, y + 78)
    y += 88
    draw_box(draw, (120, y, 400, y + 50), "这次发什么？", f, "#F3E5F5")
    y += 70

    draw_box(draw, (80, y, 360, y + 90), "发散件 ZY-D\n从【主库】出库", f, "#E8F5E9")
    draw_box(draw, (420, y, 700, y + 90), "发组合包 ZY-Z\n从【中间库】出库", f, "#E3F2FD")
    draw.text((300, y + 30), "← 选", fill="#333", font=fb)
    draw.text((380, y + 30), "选 →", fill="#333", font=fb)

    y += 120
    draw_box(draw, (80, y, 980, y + 70), "要点：组合包在商务阶段已备好；发货时只选发散件还是发组合包，不在发货时临时拼包", fs, "#FFFDE7")

    img.save(path)


def save_flowchart_warehouse(path):
    w, h = 1000, 700
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    f = get_font(17)
    fb = get_font(20, bold=True)

    draw.text((300, 15), "图2  两个仓库分别存什么", fill="#111", font=fb)

    draw_box(draw, (60, 70, 460, 320), "主库\n\nZY-D-0061 散件\nZY-D-0062 散件\nZY-D-0063 散件\nZY-D-0064 散件", f, "#E3F2FD")
    draw_box(draw, (540, 70, 940, 320), "中间库\n\nZY-Z 密封圈组合包\n（批号 / 包数 / 成本）", f, "#FFF3E0")

    arrow_right(draw, 460, 195, 540)
    draw.text((470, 160), "商务组合\n扣散件", fill="#E65100", font=f)
    draw.text((470, 210), "增组合包→", fill="#E65100", font=f)

    draw_box(draw, (60, 380, 460, 470), "销售发散件 → 客户", f, "#E8F5E9")
    draw_box(draw, (540, 380, 940, 470), "销售发组合包 → 客户", f, "#E8F5E9")
    arrow_down(draw, 260, 320, 380)
    arrow_down(draw, 740, 320, 380)

    draw_box(draw, (60, 520, 940, 620), "主库：采购入库加散件；商务组合减散件；销售可发散件\n中间库：商务组合加组合包；销售发组合包减库存", get_font(16), "#FFFDE7")
    img.save(path)


def save_flowchart_ship(path):
    w, h = 1000, 1100
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    f = get_font(17)
    fb = get_font(20, bold=True)
    fs = get_font(15)

    draw.text((280, 15), "图3  发货时如何选择（重点）", fill="#111", font=fb)

    y = 60
    draw_box(draw, (300, y, 700, y + 50), "销售创建发货单", f, "#E8F5E9")
    arrow_down(draw, 500, y + 50, y + 78)
    y += 88
    draw_box(draw, (300, y, 700, y + 50), "选择产品类型", f, "#F3E5F5")
    y += 70

    # Left branch - 散件
    draw_box(draw, (60, y, 440, y + 55), "选散件（ZY-D）", f, "#E8F5E9")
    ly = y + 65
    for t in ["查【主库】库存", "够 → 主库出库", "不够 → 禁止发货"]:
        draw_box(draw, (60, ly, 440, ly + 48), t, fs, "#F1F8E9" if "够" in t else "#FFEBEE")
        if t != "不够 → 禁止发货":
            arrow_down(draw, 250, ly + 48, ly + 62)
        ly += 62

    # Right branch - 组合包
    draw_box(draw, (560, y, 940, y + 55), "选组合包（ZY-Z）", f, "#E3F2FD")
    ry = y + 65
    for t in ["查【中间库】库存", "够 → 中间库出库", "不够 → 需商务先组合"]:
        draw_box(draw, (560, ry, 940, ry + 48), t, fs, "#E3F2FD" if "够" in t else "#FFEBEE")
        if t != "不够 → 需商务先组合":
            arrow_down(draw, 750, ry + 48, ry + 62)
        ry += 62

    draw.text((460, y + 15), "←", fill="#333", font=fb)
    draw.text((500, y + 15), "→", fill="#333", font=fb)

    y = max(ly, ry) + 30
    draw_box(draw, (60, y, 940, y + 80), "发货时选的是「产品」，不是临时拼包\n散件走原出库逻辑；组合包必须中间库已有库存", f, "#FFFDE7")
    img.save(path)


def save_flowchart_example(path):
    w, h = 1000, 900
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    f = get_font(16)
    fb = get_font(20, bold=True)

    draw.text((250, 15), "图4  示例：密封圈 60包组合 + 两种发货", fill="#111", font=fb)

    steps = [
        ("Step1 采购入库", "主库：0061×600, 0062×100, 0063×200, 0064×600\n中间库：0包"),
        ("Step2 商务组合60包", "主库扣：360/60/120/600\n中间库 +60包组合包\n剩余：240/40/80/0"),
        ("Step3A 销售发组合包10包", "选 ZY-Z 组合包 → 中间库 -10 → 剩50包"),
        ("Step3B 销售发散件0061×100", "选 ZY-D-0061 → 主库 -100 → 剩140\n（与原散件出库规则相同）"),
    ]
    y = 60
    colors = ["#E3F2FD", "#FFF3E0", "#E8F5E9", "#F3E5F5"]
    for title, body in steps:
        draw_box(draw, (60, y, 940, y + 100), f"{title}\n{body}", f, colors[steps.index((title, body))])
        if title != "Step3B 销售发散件0061×100":
            arrow_down(draw, 500, y + 100, y + 118)
        y += 118
    img.save(path)


# --- Word helpers ---

def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(11)
    run.bold = bold
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "D9E2F3")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_image(doc, path, width_inches=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_inches))
    doc.add_paragraph()


def build():
    os.makedirs(ASSET_DIR, exist_ok=True)
    img1 = os.path.join(ASSET_DIR, "flow1_overall.png")
    img2 = os.path.join(ASSET_DIR, "flow2_warehouse.png")
    img3 = os.path.join(ASSET_DIR, "flow3_ship.png")
    img4 = os.path.join(ASSET_DIR, "flow4_example.png")
    save_flowchart_overall(img1)
    save_flowchart_warehouse(img2)
    save_flowchart_ship(img3)
    save_flowchart_example(img4)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("执与自组配件出入库\n开发需求说明")
    tr.bold = True
    tr.font.size = Pt(22)
    tr.font.name = "微软雅黑"
    tr._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"版本：{DOC_VERSION}  |  日期：2026-08-21  |  用途：开发实现依据")
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_paragraph()

    # 1
    add_heading(doc, "1. 需求概述", 1)
    add_para(
        doc,
        "在现有 ERP 出入库基础上，新增「散件 → 组合包 → 销售」链路。"
        "散件（ZY-D）采购入库规则不变；新增商务维护 BOM、发起组合加工；"
        "组合包（ZY-Z）入中间库；销售发货时可选发散件（主库）或发组合包（中间库）。"
    )

    # 2 已确认
    add_heading(doc, "2. 已确认业务决策", 1)
    add_table(doc, ["项", "结论"], [
        ["BOM 维护", "商务"],
        ["组合加工发起", "商务"],
        ["组合模式", "备货型：先组合进中间库，销售再发货"],
        ["销售是否参与组合", "否，销售只在发货时选产品"],
        ["散件原入库/出库", "不变；组合扣料是新增出库类型"],
        ["发货选包时机", "销售发货时选组合包；不在发货时临时拼包"],
    ], [4, 12])

    # 3 流程图
    add_heading(doc, "3. 业务流程图", 1)
    add_image(doc, img1)
    add_image(doc, img2)
    add_image(doc, img3)
    add_image(doc, img4)

    # 4 仓库规则
    add_heading(doc, "4. 仓库与产品规则", 1)
    add_table(doc, ["仓库", "存放", "入库来源", "出库场景"], [
        ["主库", "散件 ZY-D", "采购入库（现有）", "①商务组合扣料 ②销售发散件 ③其他现有出库"],
        ["中间库", "组合包 ZY-Z", "商务组合加工完成", "销售发组合包"],
    ], [2.5, 3, 4, 6.5])

    # 5 功能
    add_heading(doc, "5. 功能需求（开发清单）", 1)
    add_table(doc, ["ID", "模块", "需求描述", "优先级"], [
        ["F1", "BOM", "维护组合包与散件用量比例；散件从产品档案选择", "P0"],
        ["F2", "组合加工", "商务选组合包+数量；系统自动算扣料与成本", "P0"],
        ["F3", "组合审核", "审核通过后：主库散件 Out、中间库组合包 In", "P0"],
        ["F4", "中间库库存", "查询组合包库存、批号、成本", "P0"],
        ["F5", "销售出库改造", "发货可选组合包，从中间库出库并校验库存", "P0"],
        ["F6", "主库库存", "现有散件库存查询不变；组合后余量仍在主库", "P0"],
        ["F7", "批号", "一次组合 N 包共用一个批号", "P1"],
        ["F8", "成本明细", "组合单展示各散件扣料批次与成本（建议 FIFO）", "P1"],
        ["F9", "拆包", "组合包拆回散件", "不做（第一期）"],
        ["F10", "财务凭证", "自动生成", "不做（第一期）"],
    ], [1.2, 2.5, 8.3, 1.5])

    # 6 业务规则
    add_heading(doc, "6. 业务规则（必须实现）", 1)
    add_table(doc, ["ID", "规则"], [
        ["R1", "散件采购入库走现有流程，入主库"],
        ["R2", "组合时从主库扣散件，组合包入中间库（不在主库存组合包）"],
        ["R3", "组合前校验主库散件库存，不足则禁止提交"],
        ["R4", "销售发组合包只能从中间库出；发散件从主库出（原逻辑）"],
        ["R5", "剩余散件不锁定，可继续组合或按原规则销售出库"],
        ["R6", "组合需审核后生效，不能保存即扣库存"],
        ["R7", "发货时选产品类型（散件/组合包），不在发货时做组合"],
        ["R8", "中间库无库存时，组合包不能发货"],
    ], [1.2, 14.8])

    # 7 权限
    add_heading(doc, "7. 角色与权限", 1)
    add_table(doc, ["角色", "权限"], [
        ["商务", "BOM 维护；组合加工发起"],
        ["审核人（待确认）", "组合加工单审核"],
        ["仓管", "散件采购入库；实物组合配合"],
        ["销售", "发货选散件或组合包；不可发起组合"],
    ], [3, 12])

    # 8 成本
    add_heading(doc, "8. 成本规则", 1)
    add_para(doc, "同一散件多次进货价不同时，需支持按规则自动计价（无需手工填单价）。", bold=True)
    add_table(doc, ["项", "说明", "状态"], [
        ["推荐方案", "FIFO 先进先出：组合时系统自动按批次扣料并汇总成本", "待业务确认"],
        ["备选", "加权平均价", "待业务确认"],
        ["不做", "每次组合手工选批次、手工填价", "—"],
    ], [3, 10, 2.5])

    # 9 待确认
    add_heading(doc, "9. 待确认项（阻塞开发细节）", 1)
    add_table(doc, ["ID", "问题", "影响"], [
        ["T1", "组合加工单审核人", "审批流配置"],
        ["T2", "成本规则：FIFO 或 加权平均", "扣料与成本算法"],
        ["T3", "中间库：独立仓库 or 同库分区", "仓库表/库存维度"],
        ["T4", "发货是否需指定组合包批号", "出库选批逻辑"],
        ["T5", "ZY-Z 编号生成方式", "产品档案规则"],
    ], [1.2, 6, 7.8])

    # 10 验收
    add_heading(doc, "10. 验收用例（UAT）", 1)
    add_table(doc, ["步骤", "操作", "预期"], [
        ["1", "采购 4 种密封圈入库", "主库有散件库存"],
        ["2", "商务维护 BOM", "1包=0061×6+0062×1+0063×2+0064×10"],
        ["3", "商务组合 60 包并审核", "主库散件减少；中间库 +60 包"],
        ["4", "销售发组合包 10 包", "中间库 -10"],
        ["5", "销售发散件 0061×100", "主库 -100（原出库逻辑）"],
    ], [1.2, 5, 8.8])

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    print(f"Size: {os.path.getsize(OUTPUT)} bytes")


if __name__ == "__main__":
    build()
